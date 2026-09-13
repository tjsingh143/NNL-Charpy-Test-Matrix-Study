"""WP0 - Acquisition check and data dictionary.

Regenerates, from the vendored raw files only:
  outputs/wp0/checksum_verification.md
  outputs/wp0/workbook_structure.md
  outputs/wp0/data_dictionary.md  and  data_dictionary.csv
  outputs/wp0/notable_facts.md
  outputs/wp0/references_summary.csv
  outputs/wp0/dataset_version_comparison.md

Run:  .venv/bin/python src/wp0_data_dictionary.py
"""
from __future__ import annotations

import collections
import re

import numpy as np
import openpyxl
import pandas as pd

from common import (DATA_SHEET, DATASETS, NEWER_RELEASE, OLDER_RELEASE, OUTPUTS, PRIMARY_DATASET, load_charpy_data,
                    load_references, md_escape, md_table, sheet_rows, verify_checksums, write_text)
from curated_metadata import COLUMN_BY_RAW, REFERENCE_NOTES

OUT = OUTPUTS / "wp0"
FERRITIC_GROUP = "Reactor pressure vessel steel"
FORMULA_CONSTANT = re.compile(r"^=\(?\s*([0-9]*\.?[0-9]+)\s*\*\s*10\s*\^\s*([0-9]+)\s*\)?$")


# ---------------------------------------------------------------------------
# Workbook structure
# ---------------------------------------------------------------------------
def inspect_workbook(xlsx):
    wb_formula = openpyxl.load_workbook(xlsx, data_only=False)
    wb_values = openpyxl.load_workbook(xlsx, data_only=True)
    info = {"sheets": [], "formulas": [], "formula_problems": []}
    for ws in wb_formula.worksheets:
        info["sheets"].append(dict(name=ws.title, state=ws.sheet_state, dimensions=ws.dimensions,
                                   max_row=ws.max_row, max_column=ws.max_column,
                                   hidden_rows=sum(1 for d in ws.row_dimensions.values() if d.hidden),
                                   hidden_columns=sum(1 for d in ws.column_dimensions.values() if d.hidden),
                                   merged_ranges=len(ws.merged_cells.ranges)))
    ws = wb_formula[DATA_SHEET]
    ws_v = wb_values[DATA_SHEET]
    header = [c.value for c in ws[2]]
    # Row-1 group headings are merged ranges; assign each column its group.
    group = {}
    for rng in ws.merged_cells.ranges:
        if rng.min_row == 1:
            for col in range(rng.min_col, rng.max_col + 1):
                group[col] = ws.cell(row=1, column=rng.min_col).value
    for col in range(1, len(header) + 1):
        if col not in group:
            group[col] = ws.cell(row=1, column=col).value or "(no heading)"
    info["group_by_header"] = {header[c - 1]: group[c] for c in range(1, len(header) + 1)}
    info["header"] = header
    # Formula cells: every one must be a pure constant whose stored (cached) result matches its evaluation.
    for row in ws.iter_rows(min_row=3):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cached = ws_v.cell(row=cell.row, column=cell.column).value
                m = FORMULA_CONSTANT.match(cell.value.replace(" ", ""))
                evaluated = float(m.group(1)) * 10 ** int(m.group(2)) if m else None
                ok = (m is not None and isinstance(cached, (int, float))
                      and np.isclose(float(cached), evaluated, rtol=1e-12, atol=0))
                rec = dict(row=cell.row, column=header[cell.column - 1], formula=cell.value, cached=cached, evaluated=evaluated, ok=ok)
                info["formulas"].append(rec)
                if not ok:
                    info["formula_problems"].append(rec)
    wb_formula.close()
    wb_values.close()
    return info


def reader_crosscheck(xlsx):
    """Compare openpyxl read-only vs normal mode, and document what pandas.read_excel does (D-003)."""
    normal = openpyxl.load_workbook(xlsx, data_only=True)[DATA_SHEET]
    rows_normal = [tuple(r) for r in normal.iter_rows(values_only=True)]
    rows_readonly = sheet_rows(xlsx, DATA_SHEET)
    n_diff = sum(1 for a, b in zip(rows_normal, rows_readonly) for x, y in zip(a, b) if x != y)
    via_pandas = pd.read_excel(xlsx, sheet_name=DATA_SHEET, header=1)
    col = "Irradiation Dose (n/cm^2)"
    return dict(rows_normal=len(rows_normal), rows_readonly=len(rows_readonly), cells_differing=n_diff,
                pandas_dtype=str(via_pandas[col].dtype),
                pandas_blank_strings=int(sum(1 for v in via_pandas[col] if isinstance(v, str) and v == "")),
                pandas_python_types=dict(collections.Counter(type(v).__name__ for v in via_pandas[col])))


# ---------------------------------------------------------------------------
# Data dictionary
# ---------------------------------------------------------------------------
def fill_category(fraction):
    # Descriptive bands only (D-026); usability for a given analysis is decided later.
    if fraction >= 0.95:
        return "complete or nearly complete"
    if fraction >= 0.50:
        return "partially populated"
    if fraction >= 0.10:
        return "sparse"
    return "very sparse (unusable for dataset-wide analysis)"


def fmt_num(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    v = float(v)
    if v != 0 and (abs(v) >= 1e6 or abs(v) < 1e-3):
        return f"{v:.4g}"
    return f"{v:.6g}"


def value_counts_text(series, numeric):
    vc = series.dropna().value_counts()
    if numeric:
        vc = vc.sort_index()
        return "; ".join(f"{fmt_num(k)} ({v})" for k, v in vc.items())
    return "; ".join(f"'{k}' ({v})" for k, v in vc.items())


def build_dictionary(df, structure, coldesc):
    ferritic = df["Material Group"] == FERRITIC_GROUP
    rows = []
    for pos, raw in enumerate(structure["header"], start=1):
        if raw not in COLUMN_BY_RAW:
            raise KeyError(f"Column {raw!r} has no entry in curated_metadata.COLUMNS")
        meta = COLUMN_BY_RAW[raw]
        s = df[raw]
        numeric = s.dtype.kind == "f"
        n_all, n_fer = int(s.notna().sum()), int(s[ferritic].notna().sum())
        row = dict(position=pos, raw_header=raw, header_group=structure["group_by_header"][raw], clean_name=meta["clean"],
                   stated_unit=meta["unit"], storage=("numeric" if numeric else "text"),
                   n_nonmissing_all=n_all, fill_all=n_all / len(df),
                   n_nonmissing_ferritic=n_fer, fill_ferritic=n_fer / int(ferritic.sum()),
                   fill_category_ferritic=fill_category(n_fer / int(ferritic.sum())),
                   n_distinct_all=int(s.nunique(dropna=True)),
                   min_all=s.min() if numeric else None, max_all=s.max() if numeric else None,
                   min_ferritic=s[ferritic].min() if numeric else None, max_ferritic=s[ferritic].max() if numeric else None,
                   meaning=meta["meaning"], authors_description=coldesc.get(meta["coldesc"], ""),
                   authors_description_name=meta["coldesc"], notes=meta["notes"])
        list_values = (not numeric) or row["n_distinct_all"] <= 25
        row["distinct_values_all"] = value_counts_text(s, numeric) if list_values else "(numeric; more than 25 distinct values - see range)"
        rows.append(row)
    return pd.DataFrame(rows)


def dictionary_markdown(dd, df, dataset_key, structure):
    n_fer = int((df["Material Group"] == FERRITIC_GROUP).sum())
    lines = [
        "# Data dictionary - Charpy impact test dataset",
        "",
        f"Dataset: {DATASETS[dataset_key]['label']}. Workbook sheet `{DATA_SHEET}`: "
        f"{len(df)} records (spreadsheet rows 3-{len(df) + 2}), {len(structure['header'])} columns.",
        "",
        "Generated by `src/wp0_data_dictionary.py` from the vendored raw file. Meanings and notes come from "
        "`src/curated_metadata.py` (sources: the workbook's own Column Descriptions sheet, the descriptor paper "
        "Lee et al. 2025, ASTM E23-82 and ISO 148-1:2016). Where a meaning is uncertain the note says so.",
        "",
        "How to read the fill columns: *all* = all records; *ferritic* = the "
        f"{n_fer} records whose Material Group is '{FERRITIC_GROUP}' (the brief's RPV-steel family; the other group, 'Nuclear Structural "
        "Material', holds the SS316 family and, in the 2025.30 release, the ferritic WWER-440 steel 15Cr2MoV-A). "
        "Fill-rate bands (D-026): >=95% complete or nearly complete; 50-95% partially populated; 10-50% sparse; "
        "<10% very sparse (unusable for dataset-wide analysis). Blank cells mean *not reported*; nothing has been filled in. "
        "Bands describe dataset-wide coverage: a field that only applies to some records (e.g. fluence, which exists only for "
        "irradiated specimens) can be complete for those records yet fall in a low band; conditional completeness is given in "
        "`notable_facts.md`.",
        "",
        "## Summary table",
        "",
    ]
    summary = dd[["position", "raw_header", "clean_name", "stated_unit", "storage", "fill_all", "fill_ferritic",
                  "fill_category_ferritic", "n_distinct_all", "min_ferritic", "max_ferritic"]].copy()
    summary["raw_header"] = summary["raw_header"].map(lambda h: f"`{h!r}`")
    summary["fill_all"] = summary["fill_all"].map(lambda v: f"{100 * v:.1f}%")
    summary["fill_ferritic"] = summary["fill_ferritic"].map(lambda v: f"{100 * v:.1f}%")
    summary["min_ferritic"] = summary["min_ferritic"].map(fmt_num)
    summary["max_ferritic"] = summary["max_ferritic"].map(fmt_num)
    summary.columns = ["#", "raw header (exact)", "clean name", "unit", "type", "fill (all)", "fill (ferritic)",
                       "fill band (ferritic)", "distinct (all)", "min (ferritic)", "max (ferritic)"]
    lines += [md_table(summary), "", "## Column details", ""]
    for r in dd.itertuples(index=False):
        lines += [f"### {r.position}. `{r.raw_header!r}` -> `{r.clean_name}`", "",
                  f"- **Group (row 1 heading):** {md_escape(r.header_group)}",
                  f"- **Unit:** {r.stated_unit or 'none'}  |  **Type:** {r.storage}",
                  f"- **Meaning:** {r.meaning}",
                  f"- **Authors' description** (Column Descriptions sheet, entry `{r.authors_description_name!r}`): "
                  f"{md_escape(r.authors_description) or '(none)'}",
                  f"- **Filled:** {r.n_nonmissing_all}/{len(df)} all records ({100 * r.fill_all:.1f}%); "
                  f"{r.n_nonmissing_ferritic}/{n_fer} ferritic records ({100 * r.fill_ferritic:.1f}%) -> {r.fill_category_ferritic}",
                  f"- **Distinct values (all records):** {r.n_distinct_all}"]
        if r.storage == "numeric":
            lines.append(f"- **Range:** all records {fmt_num(r.min_all)} to {fmt_num(r.max_all)}; "
                         f"ferritic records {fmt_num(r.min_ferritic)} to {fmt_num(r.max_ferritic)}")
        lines.append(f"- **Values with counts:** {md_escape(r.distinct_values_all) or '(all blank)'}")
        if r.notes:
            lines.append(f"- **Notes / flags:** {r.notes}")
        lines.append("")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Notable facts (every count quoted in the CP1 discussion is computed here)
# ---------------------------------------------------------------------------
def notable_facts(df, refs, structure):
    fer = df[df["Material Group"] == FERRITIC_GROUP]
    mat = df["Material "].str.strip()
    facts = []
    add = lambda k, v: facts.append((k, v))
    add("Records / columns in data sheet", f"{len(df)} / {len(structure['header'])}")
    used = sorted(df["Reference"].astype(int).unique())
    add("Reference IDs in References sheet / used by records",
        f"{len(refs)} / {len(used)} (IDs {refs.ref_id.min()}-{refs.ref_id.max()}; unused: {sorted(set(refs.ref_id) - set(used)) or 'none'})")
    add("Publication years of references", f"{refs.year.min()}-{refs.year.max()}")
    add("'Nuclear Structural Material' group equals SS316 family",
        str(bool(((df["Material Group"] == "Nuclear Structural Material") == mat.str.startswith("SS")).all())))
    add("Materials in the 'Nuclear Structural Material' group (count)",
        str(mat[df["Material Group"] == "Nuclear Structural Material"].value_counts().to_dict()))
    add("Records by Material Group", "; ".join(f"{k}: {v}" for k, v in df["Material Group"].value_counts().items()))
    by_col = collections.Counter(f["column"] for f in structure["formulas"])
    add("Formula cells (column: count) / problems", f"{dict(by_col)} / {len(structure['formula_problems'])}")
    irr = df["Irradiated Specimen"] == "Yes"
    add("Irradiated records / with fluence / with irradiation temperature",
        f"{int(irr.sum())} / {int(df.loc[irr, 'Irradiation Dose (n/cm^2)'].notna().sum())} / "
        f"{int(df.loc[irr, 'Irradiation Temperature (ºC)'].notna().sum())}")
    add("Non-irradiated records carrying a fluence value", str(int(df.loc[~irr, "Irradiation Dose (n/cm^2)"].notna().sum())))
    add("Irradiation time values", value_counts_text(df["Irradiation Time (hours)"], True))
    full = fer[(fer["Length (mm)"] == 55) & (fer["Width (mm)"] == 10) & (fer["Thickness (mm)"] == 10)]
    add("Ferritic 55x10x10 mm records missing notch depth", f"{int(full['Notch Depth (mm)'].isna().sum())} of {len(full)}")
    nodims = fer[fer[["Length (mm)", "Width (mm)", "Thickness (mm)"]].isna().all(axis=1)]
    add("Ferritic records with no specimen dimensions (by ref)", str(nodims["Reference"].astype(int).value_counts().sort_index().to_dict()))
    add("Records with blank Notch Shape (row: ref)", str({int(r): int(f) for r, f in df.loc[df["Notch Shape"].isna(), ["source_row", "Reference"]].values}))
    highp = fer[fer["P, wt.%"] > 0.1]
    add("Ferritic records with P > 0.1 wt.% (ref: count)", str(highp["Reference"].astype(int).value_counts().sort_index().to_dict()))
    add("Records with blank Test Standard", f"{int(df['Test Standard'].isna().sum())} ({100 * df['Test Standard'].isna().mean():.0f}%)")
    cid = df["DBTT Fitted Curve"]
    add("Dataset curve IDs: count, range, records with an ID", f"{cid.nunique()} IDs, {fmt_num(cid.min())}-{fmt_num(cid.max())}, {int(cid.notna().sum())} records")
    add("Records without curve ID (material: count)", str(mat[cid.isna()].value_counts().to_dict()))
    add("USE / DBTT present on records", f"{int(df['USE (J)'].notna().sum())} / {int(df['DBTT (ºC)'].notna().sum())}")
    T = df["Test Temperature (C) "]
    add("Test temperatures below absolute zero (-273.15 C): count and materials", f"{int((T < -273.15).sum())}: {mat[T < -273.15].value_counts().to_dict()}")
    add("Test temperature missing (material: count)", str(mat[T.isna()].value_counts().to_dict()))
    E = df["Impact energy (J)"]
    add("Absorbed energy: negative / exactly zero (refs)", f"{int((E < 0).sum())} / {int((E == 0).sum())} {df.loc[E == 0, 'Reference'].astype(int).tolist()}")
    add("Ferritic test temperature range (C)", f"{fmt_num(fer['Test Temperature (C) '].min())} to {fmt_num(fer['Test Temperature (C) '].max())}")
    add("Ferritic absorbed energy range (J)", f"{fmt_num(fer['Impact energy (J)'].min())} to {fmt_num(fer['Impact energy (J)'].max())}")
    frac_dec = lambda s: (s.dropna() % 1 != 0).mean()
    add("Share of non-whole-number values (digitization indicator): test temperature / energy",
        f"{100 * frac_dec(T):.0f}% / {100 * frac_dec(E):.0f}%")
    geo = fer.groupby(["Length (mm)", "Width (mm)", "Thickness (mm)"], dropna=False).size().sort_values(ascending=False)
    geo_name = lambda l, w, b: "not reported" if pd.isna(l) else f"{fmt_num(l)}x{fmt_num(w)}x{fmt_num(b)}"
    add("Ferritic specimen geometries L x W x B (count)", "; ".join(f"{geo_name(l, w, b)} ({n})" for (l, w, b), n in geo.items()))
    lines = ["# WP0 notable facts", "",
             "Computed by `src/wp0_data_dictionary.py`. 'Ferritic' = Material Group 'Reactor pressure vessel steel'.", "",
             md_table(pd.DataFrame(facts, columns=["fact", "value"]))]
    return "\n".join(lines)


def references_summary(df, refs):
    rows = []
    for r in refs.itertuples(index=False):
        s = df[df["Reference"] == r.ref_id]
        note = REFERENCE_NOTES.get(r.ref_id, {})
        rows.append(dict(ref_id=r.ref_id, year=r.year, n_records=len(s),
                         n_irradiated=int((s["Irradiated Specimen"] == "Yes").sum()),
                         materials="; ".join(f"{k} ({v})" for k, v in s["Material "].str.strip().value_counts().items()),
                         thickness_mm="; ".join(sorted({fmt_num(v) if not pd.isna(v) else "blank" for v in s["Thickness (mm)"]})),
                         n_dataset_curve_ids=int(s["DBTT Fitted Curve"].nunique()),
                         grade_in_title=note.get("grade_title", ""), product_form_note=note.get("product_form_note", ""),
                         specimen_note=note.get("specimen_note", ""), material_note=note.get("material_note", ""),
                         citation=r.citation, link=r.link))
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Version comparison (D-001)
# ---------------------------------------------------------------------------
def version_comparison(key_a, key_b):
    """Row-level comparison of an older release (A) with the newer release (B) that extends it."""
    import json
    A, B = DATASETS[key_a], DATASETS[key_b]
    role = lambda k: "primary dataset" if k == PRIMARY_DATASET else "version-sensitivity dataset"
    da, db = load_charpy_data(A["xlsx"]), load_charpy_data(B["xlsx"])
    ra, rb = load_references(A["xlsx"]), load_references(B["xlsx"])
    lines = ["# Dataset version comparison", "",
             "Generated by `src/wp0_data_dictionary.py`. Compares the older release (A, "
             f"{role(key_a)}; named in the original PROJECT_BRIEF.md) with the newer release cited by the descriptor paper "
             f"(B, {role(key_b)}). See DECISIONS.md D-001 (CP1-approved amendment).", ""]
    meta_rows = []
    for tag, key, ds, d, r in [("A", key_a, A, da, ra), ("B", key_b, B, db, rb)]:
        rec = json.loads(ds["record_json"].read_text())
        stated = re.search(r"containing ([\d,]+) test records", rec["metadata"]["description"])
        meta_rows.append(dict(release=tag, role=role(key), record=key, doi=rec["pids"]["doi"]["identifier"], issued=rec["metadata"]["publication_date"],
                              records_stated_in_archive=stated.group(1) if stated else "", records_in_file=len(d),
                              columns_in_file=d.shape[1] - 1, references_in_file=len(r),
                              sheets=", ".join(openpyxl.load_workbook(ds["xlsx"], read_only=True).sheetnames),
                              xlsx_bytes=ds["xlsx"].stat().st_size))
    lines += ["## Releases", "", md_table(pd.DataFrame(meta_rows)), ""]
    cols_a, cols_b = [c for c in da.columns if c != "source_row"], [c for c in db.columns if c != "source_row"]
    lines += ["## Columns", "",
              f"- Columns only in B: {[c for c in cols_b if c not in cols_a]}",
              f"- Columns only in A: {[c for c in cols_a if c not in cols_b]}",
              f"- Column order identical apart from additions: {[c for c in cols_b if c in cols_a] == cols_a}", ""]
    same = ra.merge(rb, on="ref_id", suffixes=("_a", "_b"))
    lines += ["## References", "",
              f"- References with identical citation text in both: {int((same.citation_a == same.citation_b).sum())} of {len(ra)} in A",
              f"- Reference IDs only in B: {sorted(set(rb.ref_id) - set(ra.ref_id))}", ""]
    for rid in sorted(set(rb.ref_id) - set(ra.ref_id)):
        n = int((db["Reference"] == rid).sum())
        mats = db.loc[db["Reference"] == rid, "Material "].str.strip().value_counts().to_dict()
        lines.append(f"  - [{rid}] {n} records {mats}: {rb.loc[rb.ref_id == rid, 'citation'].values[0]}")
    lines.append("")
    # Row-level comparison for references present in A, aligned by order within B.
    b_common = db[db["Reference"] <= ra.ref_id.max()].reset_index(drop=True)
    aligned = len(b_common) == len(da) and (b_common["Reference"].values == da["Reference"].values).all()
    lines += ["## Records of the shared references", "",
              f"- B rows belonging to A's references: {len(b_common)} (A has {len(da)}); same reference sequence row by row: {aligned}"]
    if aligned:
        changes = []
        for c in cols_a:
            a, b = da[c], b_common[c]
            if a.dtype.kind == "f" and b.dtype.kind == "f":
                diff = ~((a.isna() & b.isna()) | np.isclose(a.fillna(0), b.fillna(0), rtol=1e-12, atol=0) & (a.isna() == b.isna()))
            else:
                diff = ~((a.isna() & b.isna()) | (a.astype(str) == b.astype(str)))
            if diff.any():
                changes.append(dict(column=c, records_changed=int(diff.sum()),
                                    references=", ".join(str(int(x)) for x in sorted(da.loc[diff, "Reference"].unique()))))
        lines += [f"- Columns with at least one changed value: {len(changes)}", "", md_table(pd.DataFrame(changes)), ""]
        ia, ib = da["Irradiated Specimen"], b_common["Irradiated Specimen"]
        no_yes, yes_no = (ia == "No") & (ib == "Yes"), (ia == "Yes") & (ib == "No")
        fl_added = da["Irradiation Dose (n/cm^2)"].isna() & b_common["Irradiation Dose (n/cm^2)"].notna()
        lines += ["Irradiation status changes (A -> B):", "",
                  f"- 'No' -> 'Yes': {int(no_yes.sum())} records, refs {sorted(da.loc[no_yes, 'Reference'].astype(int).unique().tolist())}",
                  f"- 'Yes' -> 'No': {int(yes_no.sum())} records, refs {sorted(da.loc[yes_no, 'Reference'].astype(int).unique().tolist())}",
                  f"- 'No' -> 'Yes' records that also gained a fluence value in B: {int((no_yes & fl_added).sum())}",
                  f"- B records marked irradiated but without fluence: {int(((db['Irradiated Specimen'] == 'Yes') & db['Irradiation Dose (n/cm^2)'].isna()).sum())}",
                  ""]
        measurement_cols = ["Test Temperature (C) ", "Impact energy (J)", "Length (mm)", "Width (mm)", "Thickness (mm)",
                            "Notch Depth (mm)", "Notch Angle (deg)", "Notch Root Radius (mm)", "Material ", "USE (J)", "DBTT Fitted Curve"]
        changed = {ch["column"] for ch in changes}
        lines += [f"- Measurement/identity columns changed between releases: {[c for c in measurement_cols if c in changed] or 'none'}", ""]
    if "Specimen Type" in db.columns:
        st = db[db["Specimen Type"].notna()]
        lines += ["## 'Specimen Type' column (B only)", "",
                  f"- Values: {st['Specimen Type'].value_counts().to_dict()}",
                  f"- By reference (labelled / all records of that reference): "
                  + ", ".join(f"{int(r)}: {int((st['Reference'] == r).sum())}/{int((db['Reference'] == r).sum())}"
                              for r in sorted(st["Reference"].unique())), ""]
    wb = openpyxl.load_workbook(B["xlsx"], read_only=True)
    if "Notes" in wb.sheetnames:
        notes = pd.DataFrame([r for r in wb["Notes"].iter_rows(values_only=True) if any(v is not None for v in r)],
                             columns=["reference (as written)", "note (verbatim)"])
        lines += ["## 'Notes' sheet in B (verbatim)", "", md_table(notes), ""]
    return "\n".join(lines)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    checks = verify_checksums()
    write_text(OUT / "checksum_verification.md",
               "# Checksum verification\n\nAll vendored raw files match their recorded SHA-256 values, and each workbook's MD5 "
               "matches the value published in its Materials Cloud record.\n\n" + md_table(pd.DataFrame(checks)))

    ds = DATASETS[PRIMARY_DATASET]
    structure = inspect_workbook(ds["xlsx"])
    readers = reader_crosscheck(ds["xlsx"])
    df = load_charpy_data(ds["xlsx"])
    refs = load_references(ds["xlsx"])
    coldesc_rows = sheet_rows(ds["xlsx"], "Column Descriptions")
    coldesc = {name: desc for name, desc in coldesc_rows}

    header = structure["header"]
    lines = ["# Workbook structure", "", f"File: `{ds['xlsx'].relative_to(ds['xlsx'].parents[2])}` ({DATASETS[PRIMARY_DATASET]['label']})", "",
             "## Sheets", "", md_table(pd.DataFrame(structure["sheets"])), "",
             "## Data sheet layout", "",
             "- Row 1: merged group headings; row 2: column names; records start at row 3.",
             f"- Column names with leading/trailing spaces: {[h for h in header if h != h.strip()]}",
             f"- Column groups: " + "; ".join(f"{g}: {sum(1 for h in header if structure['group_by_header'][h] == g)} columns"
                                              for g in dict.fromkeys(structure["group_by_header"].values())), "",
             "## Formula cells", "",
             f"- {len(structure['formulas'])} formula cells, in columns {sorted({f['column'] for f in structure['formulas']})}.",
             f"- Every formula is a constant of the form =(a*10^b) whose stored result equals its evaluation: "
             f"{not structure['formula_problems']} ({len(structure['formula_problems'])} problems).",
             f"- Distinct formulas: {dict(collections.Counter(f['formula'] for f in structure['formulas']).most_common())}", "",
             "## Reader cross-check (D-003)", "",
             f"- openpyxl normal vs read-only mode: {readers['rows_normal']} vs {readers['rows_readonly']} rows, "
             f"{readers['cells_differing']} differing cells.",
             f"- pandas.read_excel on the fluence column: dtype `{readers['pandas_dtype']}`, Python types {readers['pandas_python_types']}, "
             f"{readers['pandas_blank_strings']} empty cells returned as the text '' instead of missing. The pipeline therefore reads "
             "cells with openpyxl and assigns types explicitly.", "",
             "## Column Descriptions sheet vs data headers", "",
             "The sheet has no header row and lists descriptions in a different order from the data columns, "
             "so descriptions are matched by name through `curated_metadata.COLUMNS` (never by position).", ""]
    cd_table = pd.DataFrame([(i + 1, f"`{n!r}`", d, header[i] if i < len(header) else "") for i, (n, d) in enumerate(coldesc_rows)],
                            columns=["row", "name in Column Descriptions", "description (verbatim)", "data header at same position"])
    lines += [md_table(cd_table)]
    unmatched = [c["coldesc"] for c in COLUMN_BY_RAW.values() if c["coldesc"] and c["raw"] in header and c["coldesc"] not in coldesc]
    lines += ["", f"- Curated description names not found in the sheet: {unmatched or 'none'}"]
    write_text(OUT / "workbook_structure.md", "\n".join(lines))

    dd = build_dictionary(df, structure, coldesc)
    dd.to_csv(OUT / "data_dictionary.csv", index=False)
    write_text(OUT / "data_dictionary.md", dictionary_markdown(dd, df, PRIMARY_DATASET, structure))
    write_text(OUT / "notable_facts.md", notable_facts(df, refs, structure))
    references_summary(df, refs).to_csv(OUT / "references_summary.csv", index=False)
    write_text(OUT / "dataset_version_comparison.md", version_comparison(OLDER_RELEASE, NEWER_RELEASE))
    print(f"WP0 complete: {len(df)} records, {len(header)} columns, {len(refs)} references -> {OUT}")


if __name__ == "__main__":
    main()
